from io import BytesIO

from openpyxl import Workbook

from bienestar.workbook import (
    _match_response_column,
    compare_workbooks,
    load_workbook,
    person_evolution,
)


INDICATORS = [
    (1, "De sí mismo", "Bienestar personal"),
    (2, "De sí mismo", "Reflexión emocional"),
    (3, "Entorno", "Red de apoyo"),
    (4, "Entorno", "Establecimiento de límites"),
]


def make_workbook(rows, *, duplicate=False, setup_suffix="") -> bytes:
    workbook = Workbook()
    setup = workbook.active
    setup.title = "SET UP"
    setup.append([None, None, None, None, "Resultado Indicador"])
    setup.append([None, "Indicador #", "Dimensión", "Nombre corto", "ROJO", "AMARILLO", "VERDE"])
    for number, dimension, label in INDICATORS:
        setup.append(
            [
                None,
                number,
                dimension,
                label,
                f"{label}: rojo{setup_suffix}",
                f"{label}: amarillo{setup_suffix}",
                f"{label}: verde{setup_suffix}",
            ]
        )

    responses = workbook.create_sheet("RESPUESTAS")
    responses.append(
        ["Número de Cédula", "Nombre Completo"]
        + [f"{label}: Por favor escoge una opción" for _, _, label in INDICATORS]
    )
    for person_id, name, scores in rows:
        answers = [f"{label}: {['', 'rojo', 'amarillo', 'verde'][score]}{setup_suffix}" for (_, _, label), score in zip(INDICATORS, scores)]
        responses.append([person_id, name, *answers])
    if duplicate:
        responses.append(responses[2])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_comparison_and_total_metrics():
    entry_bytes = make_workbook([("123", "Persona Uno", [1, 2, 3, 2])])
    exit_bytes = make_workbook([("123", "Persona Uno", [2, 1, 3, 3])])
    entry = load_workbook("entrada.xlsx", entry_bytes)
    exit_ = load_workbook("salida.xlsx", exit_bytes)

    assert entry.ok, entry.errors
    assert exit_.ok, exit_.errors
    comparison, errors = compare_workbooks(entry, exit_)
    assert not errors

    evolution, metrics = person_evolution(comparison, "123")
    assert len(evolution) == 4
    assert metrics == {"improved": 2, "worsened": 1, "unchanged": 1, "net": 1}


def test_duplicate_identifier_is_rejected():
    workbook = load_workbook(
        "duplicados.xlsx",
        make_workbook([("123", "Persona Uno", [1, 2, 3, 2])], duplicate=True),
    )
    assert not workbook.ok
    assert any("duplicadas" in error for error in workbook.errors)


def test_different_setups_cannot_be_compared():
    entry = load_workbook("entrada.xlsx", make_workbook([("123", "Persona Uno", [1, 2, 3, 2])]))
    exit_ = load_workbook(
        "salida.xlsx",
        make_workbook([("123", "Persona Uno", [1, 2, 3, 2])], setup_suffix=" cambiado"),
    )
    comparison, errors = compare_workbooks(entry, exit_)
    assert comparison is None
    assert any("SET UP" in error for error in errors)


def test_question_matching_ignores_connector_words():
    column, error = _match_response_column(
        ["Consulta a profesionales de la salud: escoge una opción"],
        "Consulta profesionales de la salud",
    )
    assert error is None
    assert column == "Consulta a profesionales de la salud: escoge una opción"


def test_setup_sheet_without_space_is_accepted():
    content = make_workbook([("123", "Persona Uno", [1, 2, 3, 2])])
    source = BytesIO(content)
    # La detección directa se cubre con un libro mínimo cuyo nombre es SETUP.
    from openpyxl import load_workbook as openpyxl_load_workbook

    loaded = openpyxl_load_workbook(source)
    loaded["SET UP"].title = "SETUP"
    output = BytesIO()
    loaded.save(output)
    parsed = load_workbook("setup_sin_espacio.xlsx", output.getvalue())
    assert parsed.ok, parsed.errors
