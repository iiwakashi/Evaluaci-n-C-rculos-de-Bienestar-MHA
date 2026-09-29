from __future__ import annotations

from io import BytesIO
import math

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from bienestar.visualization import COLORS
from bienestar.workbook import ComparisonData, person_evolution


def comparison_summary_rows(comparison: ComparisonData) -> list[dict]:
    """Construye una fila de resumen por cada persona comparable."""
    rows: list[dict] = []
    for _, person in comparison.people.iterrows():
        _, metrics = person_evolution(comparison, person["_id"])
        rows.append(
            {
                "Nombre": person["_display_name"],
                "Cédula": person["_id"],
                "Mejoran": metrics["improved"],
                "Empeoran": metrics["worsened"],
                "No cambian": metrics["unchanged"],
                "Resultado neto": metrics["net"],
            }
        )
    return rows


def comparison_summary_xlsx(comparison: ComparisonData) -> bytes:
    """Genera un Excel filtrable con una fila por persona comparable."""
    headers = ["Nombre", "Cédula", "Mejoran", "Empeoran", "No cambian", "Resultado neto"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resumen"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    sheet.append(headers)
    for row in comparison_summary_rows(comparison):
        sheet.append([row[header] for header in headers])

    header_fill = PatternFill("solid", fgColor="235347")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="182026")
            cell.alignment = Alignment(
                horizontal="left" if cell.column <= 2 else "right",
                vertical="center",
            )
        row[1].number_format = "@"

    widths = {"A": 34, "B": 18, "C": 12, "D": 12, "E": 14, "F": 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    last_row = max(sheet.max_row, 2)
    table = Table(displayName="ResumenComparables", ref=f"A1:F{last_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)

    output = BytesIO()
    workbook.save(output)
    # Reabrir verifica que la descarga sea un XLSX legible antes de entregarla.
    load_workbook(BytesIO(output.getvalue()), read_only=True).close()
    return output.getvalue()


def _wrapped_lines(text: str, max_width: float, font_size: float, max_lines: int = 3) -> list[str]:
    words = str(text).split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if not current or stringWidth(candidate, "Helvetica", font_size) <= max_width:
            current = candidate
            continue
        lines.append(current)
        current = word
        if len(lines) == max_lines - 1:
            break
    remaining_words = words[len(" ".join(lines + ([current] if current else [])).split()) :]
    if current:
        if remaining_words:
            suffix = "..."
            while current and stringWidth(current + suffix, "Helvetica", font_size) > max_width:
                current = current[:-1]
            current += suffix
        lines.append(current)
    return lines[:max_lines]


def _draw_centered_lines(
    pdf: canvas.Canvas,
    lines: list[str],
    center_x: float,
    start_y: float,
    font_size: float,
) -> None:
    pdf.setFont("Helvetica", font_size)
    pdf.setFillColor(HexColor("#182026"))
    for index, line in enumerate(lines):
        pdf.drawCentredString(center_x, start_y - index * (font_size + 1.5), line)


def person_view_pdf(name: str, evolution: list[dict], metrics: dict[str, int]) -> bytes:
    """Dibuja la vista individual en una página carta horizontal."""
    output = BytesIO()
    page_width, page_height = landscape(letter)
    pdf = canvas.Canvas(output, pagesize=(page_width, page_height))
    margin = 30

    pdf.setTitle(f"Evolución - {name}")
    pdf.setFillColor(HexColor("#173F35"))
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(margin, page_height - 36, "Evolución de Círculos de Bienestar")
    pdf.setFillColor(HexColor("#182026"))
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawString(margin, page_height - 58, name)

    legend_y = page_height - 79
    for index, (label, color) in enumerate(
        (("Rojo", COLORS[1]), ("Amarillo", COLORS[2]), ("Verde", COLORS[3]))
    ):
        x = margin + index * 85
        pdf.setFillColor(HexColor(color))
        pdf.circle(x + 6, legend_y + 2, 6, fill=1, stroke=0)
        pdf.setFillColor(HexColor("#38444D"))
        pdf.setFont("Helvetica", 8.5)
        pdf.drawString(x + 16, legend_y - 1, label)
    pdf.drawRightString(
        page_width - margin,
        legend_y - 1,
        "Círculo pequeño: entrada   |   Círculo grande: salida",
    )

    count = max(len(evolution), 1)
    columns = min(6, count)
    rows = math.ceil(count / columns)
    grid_top = legend_y - 18
    summary_top = 88
    grid_height = grid_top - summary_top
    cell_width = (page_width - 2 * margin) / columns
    cell_height = grid_height / rows
    label_font = max(6.5, min(9.0, cell_height / 10))
    exit_radius = max(11, min(18, cell_height * 0.18))
    entry_radius = exit_radius * 0.62

    for index, item in enumerate(evolution):
        row = index // columns
        column = index % columns
        center_x = margin + column * cell_width + cell_width / 2
        center_y = grid_top - row * cell_height - cell_height * 0.34

        pdf.setFillColor(HexColor(COLORS[item["entry"]]))
        pdf.circle(center_x - exit_radius * 0.65, center_y, entry_radius, fill=1, stroke=0)
        pdf.setStrokeColor(white)
        pdf.setLineWidth(1.5)
        pdf.setFillColor(HexColor(COLORS[item["exit"]]))
        pdf.circle(center_x + exit_radius * 0.35, center_y, exit_radius, fill=1, stroke=1)

        lines = _wrapped_lines(item["label"], cell_width - 12, label_font)
        _draw_centered_lines(pdf, lines, center_x, center_y - exit_radius - 13, label_font)

    pdf.setStrokeColor(HexColor("#D9DDE1"))
    pdf.line(margin, 82, page_width - margin, 82)
    pdf.setFillColor(HexColor("#182026"))
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(margin, 68, "Resumen total de evolución")

    metric_items = [
        ("Mejoran", metrics["improved"], "#2FA34A"),
        ("Empeoran", metrics["worsened"], "#E94F4F"),
        ("Siguen igual", metrics["unchanged"], "#7A838B"),
        (
            "Resultado neto",
            f"+{metrics['net']}" if metrics["net"] > 0 else metrics["net"],
            "#157A3D" if metrics["net"] > 0 else "#B92727" if metrics["net"] < 0 else "#46515A",
        ),
    ]
    gap = 8
    card_width = (page_width - 2 * margin - gap * 3) / 4
    for index, (label, value, color) in enumerate(metric_items):
        x = margin + index * (card_width + gap)
        pdf.setFillColor(HexColor(color))
        pdf.roundRect(x, 26, card_width, 32, 8, fill=1, stroke=0)
        pdf.setFillColor(white)
        pdf.setFont("Helvetica-Bold", 8.5)
        pdf.drawString(x + 10, 38, label)
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawRightString(x + card_width - 10, 36, str(value))

    pdf.showPage()
    pdf.save()
    return output.getvalue()
